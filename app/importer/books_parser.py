"""Lector del Excel de catálogo inicial de libros (BE-17).

El fichero de importación es la hoja de cálculos del centro (más de cien
títulos mantenidos a mano durante años). El formato canónico acordado es una
hoja única con las columnas:

    titulo;nivel;ejemplares;sesiones

`ejemplares` y `sesiones` son texto libre (escenario 4): valores como
`11 Fotocopias`, `PDF`, `PREPARAR` o `5+` no rechazan la fila, se guardan en
`copies_note` y `sessions_note`. Un valor numérico de fecha de Excel (p. ej.
la serie `43075`) se convierte a texto.

La robustez del apartado de notas de la historia se cubre así:
- cabeceras repetidas a mitad de hoja se ignoran (escenario 2);
- los rótulos de sección en mayúsculas (`ELIGE TU PROPIA AVENTURA`) sin nivel
  se reconocen como separadores y no se cargan (escenario 3);
- las filas de las que no se puede extraer título o nivel van a un listado de
  revisión manual y el resto del fichero continúa (escenario 5).
"""
import csv
import datetime
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from app.core.exceptions import ValidationError

MAX_TITLE_LENGTH = 255
MAX_NOTE_LENGTH = 255

TITLE_HEADERS = {
    "titulo",
    "título",
    "titulo del libro",
    "título del libro",
    "libro",
    "title",
}
LEVEL_HEADERS = {"nivel", "nivel del libro", "level"}
COPIES_HEADERS = {
    "ejemplares",
    "ejemplar",
    "copias",
    "copies",
    "copies_note",
    "nº ejemplares",
    "no ejemplares",
    "n.º de ejemplares",
    "numero de ejemplares",
    "centro y nº libros",
    "centro y n.º libros",
}
SESSIONS_HEADERS = {"sesiones", "sesión", "sessions", "sessions_note"}

# Niveles canónicos BE-15 con sus variantes de escritura habituales.
CANONICAL_BY_ALIAS = {
    "0": "0",
    "0-i": "0-I",
    "0i": "0-I",
    "i": "I",
    "ii": "II",
    "i-ii": "I/II",
    "i/ii": "I/II",
}

REVIEW_HEADER = ["linea", "titulo", "nivel", "ejemplares", "sesiones", "motivo"]


def _cell_key(value: Any) -> str:
    """Clave normalizada de una celda para comparar cabeceras."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def normalize_book_level(value: Any) -> Optional[str]:
    """Normaliza el texto de la hoja al nivel canónico BE-15.

    Acepta variantes como "Nivel 0-I", "N. 0-I", "0 - I", "I-II" y devuelve
    uno de los cinco valores de `BookLevel`; None si no se reconoce.
    """
    if value is None:
        return None

    key = re.sub(r"\s+", "", str(value).strip().lower())
    key = re.sub(r"^nivel", "", key)
    key = re.sub(r"^n\.?", "", key)
    return CANONICAL_BY_ALIAS.get(key, None)


def _cell_to_text(value: Any) -> Optional[str]:
    """Convierte una celda (texto, número, fecha de Excel) a texto plano."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _truncate(value: Optional[str], limit: int) -> Optional[str]:
    if value is None:
        return None
    return value[:limit]


def _is_title_header(value: Any) -> bool:
    return _cell_key(value) in TITLE_HEADERS


def _is_section_label(title: Any) -> bool:
    """Rótulo de sección: título en mayúsculas sin un solo carácter en minúscula."""
    if title is None:
        return False
    text = str(title).strip()
    if not text:
        return False
    has_letter = any(ch.isalpha() for ch in text)
    has_lowercase = any(ch.islower() for ch in text)
    return has_letter and not has_lowercase


def _resolve_columns(header: Tuple[Any, ...]) -> Dict[str, int]:
    cols: Dict[str, int] = {}
    for index, token in enumerate(header):
        key = _cell_key(token)
        if key in cols:
            continue
        if key in TITLE_HEADERS:
            cols["titulo"] = index
        elif key in LEVEL_HEADERS:
            cols["nivel"] = index
        elif key in COPIES_HEADERS:
            cols["ejemplares"] = index
        elif key in SESSIONS_HEADERS:
            cols["sesiones"] = index
    return cols


def _locate_header(sheet_rows: List[Tuple[Any, ...]]) -> Optional[int]:
    for row_index, row in enumerate(sheet_rows):
        if any(_is_title_header(cell) for cell in row):
            return row_index
    return None


def _row_is_empty(row: Tuple[Any, ...]) -> bool:
    return all(not _cell_key(cell) for cell in row)


def _raw_text(value: Any, limit: int = MAX_NOTE_LENGTH) -> Optional[str]:
    return _truncate(_cell_to_text(value), limit)


def parse_books_xlsx(source) -> Dict[str, Any]:
    """Parsea el Excel del catálogo y devuelve `rows`, `review` y contadores.

    Resultado:
        rows: listas de dicts {titulo, nivel, copies_note, sessions_note, line}
        review: filas no interpretables {line, titulo, nivel, ejemplares,
                sesiones, motivo} para revisión manual (escenario 5)
        headers_skipped / separators_skipped: contadores de filas ignoradas

    Levanta ValidationError si ninguna hoja tiene cabecera de título o si la
    cabecera localizada no incluye la columna de nivel.
    """
    workbook = load_workbook(str(source), data_only=True)

    sheet_rows = None
    cols: Dict[str, int] = {}
    header_row_index: Optional[int] = None

    for worksheet in workbook.worksheets:
        rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        candidate = _locate_header(rows)
        if candidate is None:
            continue
        resolved = _resolve_columns(rows[candidate])
        if "titulo" not in resolved or "nivel" not in resolved:
            continue
        sheet_rows = rows
        cols = resolved
        header_row_index = candidate
        break

    if sheet_rows is None or header_row_index is None:
        raise ValidationError(
            "El archivo Excel debe contener una hoja con las columnas "
            "titulo(nombre del libro) y nivel"
        )

    rows: List[Dict[str, Any]] = []
    review: List[Dict[str, Any]] = []
    headers_skipped = 0
    separators_skipped = 0

    for index in range(header_row_index + 1, len(sheet_rows)):
        line = index + 1
        row = sheet_rows[index]

        if _row_is_empty(row):
            continue

        title_raw = row[cols["titulo"]]
        level_raw = row[cols["nivel"]]
        nivel = normalize_book_level(level_raw)

        if _is_title_header(title_raw):
            headers_skipped += 1
            continue

        title_text = _truncate(_cell_to_text(title_raw), MAX_TITLE_LENGTH)

        if _is_section_label(title_raw) and nivel is None:
            separators_skipped += 1
            continue

        if title_text is None:
            review.append(
                _review_entry(line, row, cols, title_raw, level_raw, "No se puede extraer el título")
            )
            continue

        if nivel is None:
            raw_level = _cell_to_text(level_raw) or ""
            if raw_level:
                motivo = f"Nivel '{raw_level}' no reconocido"
            else:
                motivo = "No se puede extraer el nivel"
            review.append(_review_entry(line, row, cols, title_raw, level_raw, motivo))
            continue

        rows.append({
            "titulo": title_text,
            "nivel": nivel,
            "copies_note": _raw_text(row[cols["ejemplares"]] if "ejemplares" in cols else None),
            "sessions_note": _raw_text(row[cols["sesiones"]] if "sesiones" in cols else None),
            "line": line,
        })

    return {
        "rows": rows,
        "review": review,
        "headers_skipped": headers_skipped,
        "separators_skipped": separators_skipped,
    }


def _review_entry(line, row, cols, title_raw, level_raw, motivo: str) -> Dict[str, Any]:
    return {
        "line": line,
        "titulo": _raw_text(title_raw, MAX_TITLE_LENGTH) or "",
        "nivel": _raw_text(level_raw) or "",
        "ejemplares": _raw_text(row[cols["ejemplares"]] if "ejemplares" in cols else None) or "",
        "sesiones": _raw_text(row[cols["sesiones"]] if "sesiones" in cols else None) or "",
        "motivo": motivo,
    }


def write_review_csv(source, review: List[Dict[str, Any]]) -> Path:
    """Escribe el listado de revisión manual junto al fichero importado.

    Genera `<nombre>_revision.csv` (separador ';') con las filas no
    interpretables y su motivo, para que el administrador las repase a mano
    (BE-17 escenario 5 / T-BE17-05).
    """
    source_path = Path(source)
    out_path = source_path.parent / f"{source_path.stem}_revision.csv"

    with out_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(REVIEW_HEADER)
        for entry in review:
            writer.writerow(
                [
                    entry.get("line", ""),
                    entry.get("titulo", ""),
                    entry.get("nivel", ""),
                    entry.get("ejemplares", ""),
                    entry.get("sesiones", ""),
                    entry.get("motivo", ""),
                ]
            )
    return out_path