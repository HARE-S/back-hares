import datetime
import io
import uuid
import openpyxl
import pytest
from app import create_app
from app.config import Config
from app.core.audit import clear_audit_logs, get_audit_logs
from app.models.center import Center, Section
from app.models.student import Student, StudentSection
from app.models.test import Result, Test


class NoBypassConfig(Config):
    TESTING = True
    APP_ENV = "development"
    DEV_AUTH_BYPASS = False
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"


@pytest.fixture
def setup_data(session):
    """Fixture para crear datos base de prueba para BE-37."""
    clear_audit_logs()

    center1 = Center(name="Colegio Cervantes")
    center_empty = Center(name="Colegio Sin Datos")
    session.add_all([center1, center_empty])
    session.flush()

    sec1 = Section(name="1º ESO A", center_id=center1.id, academic_year="2025-2026")
    sec2 = Section(name="1º ESO B", center_id=center1.id, academic_year="2025-2026")
    sec_empty = Section(name="1º ESO C", center_id=center1.id, academic_year="2025-2026")
    session.add_all([sec1, sec2, sec_empty])
    session.flush()

    s1 = Student(name="Aitor Ortiz", external_id="AIT-001", birth_date=datetime.date(2012, 5, 15))
    s2 = Student(name="Leire Blanco", external_id="LEI-002", birth_date=datetime.date(2013, 3, 20))
    s3 = Student(name="Lucía Morales", external_id="LUC-003", birth_date=datetime.date(2012, 8, 20))
    session.add_all([s1, s2, s3])
    session.flush()

    session.add_all([
        StudentSection(student_id=s1.id, section_id=sec1.id),
        StudentSection(student_id=s2.id, section_id=sec1.id),
        StudentSection(student_id=s3.id, section_id=sec2.id),
    ])


    test1 = Test(code="1CL", name="Prueba Inicial 1", words=300)
    test2 = Test(code="2CL", name="Prueba Trimestre 2", words=300)
    session.add_all([test1, test2])
    session.flush()

    # Resultados en Sección 1:
    # r1: 300 words / 120s = 150.0 PPM; 15 aciertos, 3 fallos (P=13.5 -> CL=67.5%)
    r1 = Result(
        student_id=s1.id,
        section_id=sec1.id,
        test_id=test1.id,
        test_date=datetime.date(2025, 10, 1),
        time=120,
        successes=15,
        mistakes=3,
    )
    # r2: 300 words / 100s = 180.0 PPM; 18 aciertos, 2 fallos (P=17.0 -> CL=85.0%)
    r2 = Result(
        student_id=s1.id,
        section_id=sec1.id,
        test_id=test2.id,
        test_date=datetime.date(2025, 12, 15),
        time=100,
        successes=18,
        mistakes=2,
    )
    # r3: 300 words / 200s = 90.0 PPM; 8 aciertos, 6 fallos (P=5.0 -> CL=25.0%)
    r3 = Result(
        student_id=s2.id,
        section_id=sec1.id,
        test_id=test1.id,
        test_date=datetime.date(2025, 10, 1),
        time=200,
        successes=8,
        mistakes=6,
    )

    # Resultado en Sección 2:
    # r4: 300 words / 150s = 120.0 PPM; 16 aciertos, 0 fallos (P=16.0 -> CL=80.0%)
    r4 = Result(
        student_id=s3.id,
        section_id=sec2.id,
        test_id=test1.id,
        test_date=datetime.date(2025, 10, 1),
        time=150,
        successes=16,
        mistakes=0,
    )

    session.add_all([r1, r2, r3, r4])
    session.commit()

    return {
        "center1": center1,
        "center_empty": center_empty,
        "sec1": sec1,
        "sec2": sec2,
        "sec_empty": sec_empty,
        "s1": s1,
        "s2": s2,
        "s3": s3,
    }


def test_scenario_1_section_aggregated_metrics(client, setup_data):
    """
    Escenario 1: Datos agregados del grupo.
    Media de PPM, porcentaje de aciertos, número de participantes y pruebas realizadas.
    """
    sec1 = setup_data["sec1"]

    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["user_role"] = "coordinator"

    resp = client.get(f"/api/v1/sections/{sec1.id}/report")
    assert resp.status_code == 200

    data = resp.get_json()
    assert data["section_id"] == str(sec1.id)
    assert data["name"] == "1º ESO A"
    assert data["has_data"] is True
    assert data["participants_count"] == 2
    assert data["results_count"] == 3

    # Medias calculadas:
    # PPMs: [150.0, 180.0, 90.0] -> sum = 420.0 / 3 = 140.0
    # Accuracy: [67.5, 85.0, 25.0] -> sum = 177.5 / 3 = 59.17
    assert data["mean_ppm"] == 140.0
    assert pytest.approx(data["mean_accuracy"], 0.01) == 59.17
    assert data["mean_vef"] is not None
    assert "generated_at" in data


def test_scenario_2_distribution_not_only_mean(client, setup_data):
    """
    Escenario 2: Distribución, no solo media.
    Verifica que incluye la distribución de resultados por tramos y estadísticos.
    """
    sec1 = setup_data["sec1"]

    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["user_role"] = "coordinator"

    resp = client.get(f"/api/v1/sections/{sec1.id}/report")
    assert resp.status_code == 200
    data = resp.get_json()

    dist = data.get("distribution")
    assert dist is not None

    # Tramos de Comprensión
    acc_brackets = {item["bracket"]: item for item in dist["accuracy_brackets"]}
    assert acc_brackets["<50"]["count"] == 1  # 25.0%
    assert acc_brackets["50-69"]["count"] == 1  # 67.5%
    assert acc_brackets["70-84"]["count"] == 0
    assert acc_brackets[">=85"]["count"] == 1  # 85.0%
    assert pytest.approx(acc_brackets["<50"]["percentage"], 0.01) == 33.33

    # Tramos de Velocidad PPM
    ppm_brackets = {item["bracket"]: item for item in dist["ppm_brackets"]}
    assert ppm_brackets["<100"]["count"] == 1  # 90.0
    assert ppm_brackets["100-149"]["count"] == 0
    assert ppm_brackets["150-199"]["count"] == 2  # 150.0, 180.0
    assert ppm_brackets[">=200"]["count"] == 0

    # Estadísticos descriptivos
    stats = dist["stats"]
    assert "ppm" in stats
    assert "accuracy" in stats
    assert stats["ppm"]["min"] == 90.0
    assert stats["ppm"]["max"] == 180.0
    assert stats["ppm"]["median"] == 150.0


def test_scenario_3_center_report_with_section_breakdown(client, setup_data):
    """
    Escenario 3: Informe de centro.
    Devuelve datos del conjunto global y desglose por sección.
    """
    center1 = setup_data["center1"]

    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["user_role"] = "coordinator"

    resp = client.get(f"/api/v1/centers/{center1.id}/report")
    assert resp.status_code == 200
    data = resp.get_json()

    assert data["center_id"] == str(center1.id)
    assert data["name"] == "Colegio Cervantes"
    assert data["has_data"] is True
    # 3 alumnos en total (2 en sec1 + 1 en sec2)
    assert data["participants_count"] == 3
    # 4 pruebas en total (3 en sec1 + 1 en sec2)
    assert data["results_count"] == 4

    # Desglose por sección
    sections = data.get("sections")
    assert isinstance(sections, list)
    assert len(sections) == 3

    sec_map = {s["name"]: s for s in sections}
    assert "1º ESO A" in sec_map
    assert "1º ESO B" in sec_map
    assert "1º ESO C" in sec_map

    assert sec_map["1º ESO A"]["participants_count"] == 2
    assert sec_map["1º ESO A"]["results_count"] == 3
    assert sec_map["1º ESO A"]["has_data"] is True

    assert sec_map["1º ESO B"]["participants_count"] == 1
    assert sec_map["1º ESO B"]["results_count"] == 1
    assert sec_map["1º ESO B"]["mean_ppm"] == 120.0

    # Sección C no tiene datos
    assert sec_map["1º ESO C"]["has_data"] is False
    assert sec_map["1º ESO C"]["mean_ppm"] is None


def test_scenario_4_group_and_center_without_data(client, setup_data):
    """
    Escenario 4: Grupo sin datos.
    Indica ausencia de datos y no devuelve medias calculadas sobre cero.
    """
    sec_empty = setup_data["sec_empty"]
    center_empty = setup_data["center_empty"]

    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["user_role"] = "coordinator"

    # Sección vacía
    resp_sec = client.get(f"/api/v1/sections/{sec_empty.id}/report")
    assert resp_sec.status_code == 200
    data_sec = resp_sec.get_json()

    assert data_sec["has_data"] is False
    assert data_sec["participants_count"] == 0
    assert data_sec["results_count"] == 0
    assert data_sec["mean_ppm"] is None
    assert data_sec["mean_accuracy"] is None
    assert data_sec["distribution"] is None

    # Centro vacío
    resp_ctr = client.get(f"/api/v1/centers/{center_empty.id}/report")
    assert resp_ctr.status_code == 200
    data_ctr = resp_ctr.get_json()

    assert data_ctr["has_data"] is False
    assert data_ctr["participants_count"] == 0
    assert data_ctr["results_count"] == 0
    assert data_ctr["mean_ppm"] is None
    assert data_ctr["mean_accuracy"] is None


def test_scenario_5_excel_export(client, setup_data):
    """
    Escenario 5: Exportable a Excel con tipos nativos y hojas estructuradas.
    """
    sec1 = setup_data["sec1"]
    center1 = setup_data["center1"]

    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["user_role"] = "coordinator"

    # 1. Exportación de sección mediante ?format=excel
    resp_sec = client.get(f"/api/v1/sections/{sec1.id}/report?format=excel")
    assert resp_sec.status_code == 200
    assert "spreadsheetml.sheet" in resp_sec.content_type

    wb_sec = openpyxl.load_workbook(io.BytesIO(resp_sec.data))
    assert "Resumen" in wb_sec.sheetnames
    assert "Distribución" in wb_sec.sheetnames

    # Verificar que los números se guardan como tipos nativos reales (float / int)
    ws_res = wb_sec["Resumen"]
    # Buscar celda de media PPM
    found_ppm = False
    for row in ws_res.iter_rows(values_only=True):
        if row[0] == "Media Palabras Por Minuto (PPM)":
            assert isinstance(row[1], (float, int))
            assert row[1] == 140.0
            found_ppm = True
    assert found_ppm is True

    # 2. Exportación directa de sección /report/export
    resp_sec_export = client.get(f"/api/v1/sections/{sec1.id}/report/export")
    assert resp_sec_export.status_code == 200
    assert "spreadsheetml.sheet" in resp_sec_export.content_type

    # 3. Exportación de centro con desglose
    resp_ctr = client.get(f"/api/v1/centers/{center1.id}/report?format=excel")
    assert resp_ctr.status_code == 200
    assert "spreadsheetml.sheet" in resp_ctr.content_type

    wb_ctr = openpyxl.load_workbook(io.BytesIO(resp_ctr.data))
    assert "Resumen" in wb_ctr.sheetnames
    assert "Distribución" in wb_ctr.sheetnames
    assert "Desglose Secciones" in wb_ctr.sheetnames


def test_scenario_6_role_access_control(client, setup_data):
    """
    Escenario 6: Alcance por rol.
    Tutor denegado (403) para informe de centro completo.
    Tutor solo accede a secciones que tiene asignadas.
    """
    center1_id = str(setup_data["center1"].id)
    sec1_id = str(setup_data["sec1"].id)
    sec2_id = str(setup_data["sec2"].id)

    # 1. Petición sin autenticación -> 401 Unauthorized
    no_bypass_app = create_app(NoBypassConfig)
    with no_bypass_app.test_client() as unauth_client:
        assert unauth_client.get(f"/api/v1/centers/{center1_id}/report").status_code == 401
        assert unauth_client.get(f"/api/v1/sections/{sec1_id}/report").status_code == 401

    # 2. Tutor solicita informe de centro completo -> 403 Forbidden (Escenario 6 estricto)
    client.post("/api/dev/session", json={"role": "tutor", "sections": [sec1_id]})

    resp_ctr_tutor = client.get(f"/api/v1/centers/{center1_id}/report")
    assert resp_ctr_tutor.status_code == 403
    assert resp_ctr_tutor.get_json()["error"] == "FORBIDDEN"

    # 3. Tutor solicita su sección asignada (sec1) -> 200 OK
    resp_tutor_sec1 = client.get(f"/api/v1/sections/{sec1_id}/report")
    assert resp_tutor_sec1.status_code == 200

    # 4. Tutor solicita sección NO asignada (sec2) -> 403 Forbidden
    resp_tutor_sec2 = client.get(f"/api/v1/sections/{sec2_id}/report")
    assert resp_tutor_sec2.status_code == 403
    assert resp_tutor_sec2.get_json()["error"] == "FORBIDDEN"

    # 5. Usuario con rol 'pendiente' -> 403 Forbidden en ambos
    client.post("/api/dev/session", json={"role": "pendiente", "sections": [sec1_id]})
    assert client.get(f"/api/v1/sections/{sec1_id}/report").status_code == 403
    assert client.get(f"/api/v1/centers/{center1_id}/report").status_code == 403

    # 6. Coordinador accede a ambos sin restricción
    client.post("/api/dev/session", json={"role": "coordinator", "sections": []})
    assert client.get(f"/api/v1/sections/{sec1_id}/report").status_code == 200
    assert client.get(f"/api/v1/centers/{center1_id}/report").status_code == 200




def test_audit_logging_and_validation(client, setup_data):
    """
    Verifica que se registran eventos de auditoría y validación de 400 y 404.
    """
    sec1 = setup_data["sec1"]
    center1 = setup_data["center1"]

    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["user_role"] = "coordinator"

    clear_audit_logs()

    # Generar informes
    client.get(f"/api/v1/sections/{sec1.id}/report")
    client.get(f"/api/v1/centers/{center1.id}/report")

    logs = get_audit_logs()
    actions = [l["action"] for l in logs]
    assert "GENERATE_SECTION_REPORT" in actions
    assert "GENERATE_CENTER_REPORT" in actions

    # 400 Bad Request por ID inválido
    resp_invalid_sec = client.get("/api/v1/sections/no-es-uuid/report")
    assert resp_invalid_sec.status_code == 400

    resp_invalid_ctr = client.get("/api/v1/centers/no-es-uuid/report")
    assert resp_invalid_ctr.status_code == 400

    # 404 Not Found por UUID inexistente
    random_uuid = str(uuid.uuid4())
    resp_404_sec = client.get(f"/api/v1/sections/{random_uuid}/report")
    assert resp_404_sec.status_code == 404

    resp_404_ctr = client.get(f"/api/v1/centers/{random_uuid}/report")
    assert resp_404_ctr.status_code == 404
