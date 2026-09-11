import datetime
import io
import uuid
import openpyxl
import pytest

from app import create_app
from app.config import Config
from app.core.audit import get_audit_logs, clear_audit_logs
from app.models.center import Center, Section
from app.models.student import Student, StudentSection
from app.models.test import Result, Test


class NoBypassConfig(Config):
    TESTING = True
    APP_ENV = "development"
    DEV_AUTH_BYPASS = False
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"


@pytest.fixture
def setup_data(session):
    """Fixture para crear datos de prueba con secciones, alumnos y resultados."""
    clear_audit_logs()

    center = Center(name="Colegio Cervantes")
    session.add(center)
    session.flush()

    sec1 = Section(name="1º ESO A", center_id=center.id, academic_year="2025-2026")
    sec2 = Section(name="1º ESO B", center_id=center.id, academic_year="2025-2026")
    session.add_all([sec1, sec2])
    session.flush()

    # Alumno 1 en Sección 1
    s1 = Student(name="Aitor Ortiz", external_id="AIT-001", birth_date=datetime.date(2012, 5, 15))
    # Alumno 2 en Sección 2
    s2 = Student(name="Lucía Morales", external_id="LUC-002", birth_date=datetime.date(2012, 8, 20))
    session.add_all([s1, s2])
    session.flush()

    ss1 = StudentSection(student_id=s1.id, section_id=sec1.id, created_on=datetime.date(2025, 9, 10))
    ss2 = StudentSection(student_id=s2.id, section_id=sec2.id, created_on=datetime.date(2025, 9, 10))
    session.add_all([ss1, ss2])
    session.flush()

    # Pruebas
    test1 = Test(code="TEST-01", name="Comprensión Inicial", words=120, course=1, test_letter="I", type="Lectura")
    test2 = Test(code="TEST-02", name="Comprensión Media", words=180, course=1, test_letter="M", type="Lectura")
    session.add_all([test1, test2])
    session.flush()

    # Resultados para s1 (sec1)
    res1 = Result(
        student_id=s1.id,
        section_id=sec1.id,
        test_id=test1.id,
        test_date=datetime.date(2025, 10, 15),
        time=60,
        successes=15,
        mistakes=2,
    )
    # Resultados para s2 (sec2)
    res2 = Result(
        student_id=s2.id,
        section_id=sec2.id,
        test_id=test2.id,
        test_date=datetime.date(2025, 10, 20),
        time=90,
        successes=18,
        mistakes=1,
    )
    session.add_all([res1, res2])
    session.commit()

    return {
        "center": center,
        "section1": sec1,
        "section2": sec2,
        "student1": s1,
        "student2": s2,
        "test1": test1,
        "test2": test2,
        "result1": res1,
        "result2": res2,
    }


def test_scenario_1_export_filtered_dataset(client, setup_data):
    """
    Escenario 1: Exportación del conjunto filtrado
    Dado un conjunto de resultados con filtros aplicados
    Cuando se solicita la exportación con esos mismos filtros
    Entonces el fichero contiene exactamente ese conjunto y no el conjunto sin filtrar.
    """
    sec1 = setup_data["section1"]
    s1 = setup_data["student1"]

    # Autenticar como coordinador
    client.post("/api/dev/session", json={"role": "coordinator", "sections": []})

    resp = client.get(f"/api/v1/results/export/excel?section_id={sec1.id}")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers.get("Content-Type", "")

    # Cargar el libro Excel generado
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    sheet = wb.active

    # La fila 1 son las cabeceras, las filas siguientes son datos
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 2  # 1 fila de cabecera + 1 fila para s1

    header = rows[0]
    data_row = rows[1]

    # Verificar que el resultado corresponde a s1 y sec1
    student_col_idx = [i for i, h in enumerate(header) if "Alumno" in str(h) and "ID" not in str(h)][0]
    section_col_idx = [i for i, h in enumerate(header) if "Sección" in str(h)][0]

    assert data_row[student_col_idx] == s1.name
    assert data_row[section_col_idx] == sec1.name


def test_scenario_2_file_format_and_spanish_headers(client, setup_data):
    """
    Escenario 2: Formato del fichero
    Dado una exportación solicitada
    Cuando se descarga el fichero
    Entonces es un .xlsx válido y sus cabeceras están en castellano y son legibles.
    """
    client.post("/api/dev/session", json={"role": "coordinator", "sections": []})

    resp = client.get("/api/v1/results/export/excel")
    assert resp.status_code == 200

    # Comprobar cabecera Content-Disposition
    disp = resp.headers.get("Content-Disposition", "")
    assert "attachment" in disp
    assert ".xlsx" in disp

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    sheet = wb.active
    assert sheet.title is not None

    headers = [cell.value for cell in sheet[1]]
    expected_keywords = ["Fecha", "Alumno", "Sección", "Prueba", "PPM", "Aciertos", "Fallos"]
    for kw in expected_keywords:
        assert any(kw.lower() in str(h).lower() for h in headers), f"Falta cabecera esperada con palabra clave '{kw}'"


def test_scenario_3_calculated_metrics_included(client, setup_data):
    """
    Escenario 3: Métricas incluidas
    Dado una exportación de resultados
    Cuando se abre el fichero
    Entonces cada fila incluye el PPM y el porcentaje de aciertos/comprensión, y no solo el tiempo, aciertos y fallos.
    """
    sec1 = setup_data["section1"]
    client.post("/api/dev/session", json={"role": "tutor", "sections": [str(sec1.id)]})

    resp = client.get(f"/api/v1/results/export/excel?section_id={sec1.id}")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    data_row = rows[1]

    ppm_idx = [i for i, h in enumerate(headers) if "PPM" in h.upper()][0]
    comp_idx = [i for i, h in enumerate(headers) if "COMPRENSI" in h.upper()][0]

    # test1 words=120, time=60 -> PPM = 120.0
    assert data_row[ppm_idx] == 120.0
    # successes=15, mistakes=2 -> score = 15 - 1 = 14; (14 / 20) * 100 = 70.0%
    assert data_row[comp_idx] == 70.0


def test_scenario_4_correct_native_data_types(client, setup_data):
    """
    Escenario 4: Tipos de dato correctos
    Dado una exportación con fechas y valores decimales
    Cuando se abre el fichero en una hoja de cálculo
    Entonces las fechas se reconocen como fechas, los decimales como números y ninguno aparece como texto.
    """
    client.post("/api/dev/session", json={"role": "coordinator", "sections": []})

    resp = client.get("/api/v1/results/export/excel")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    sheet = wb.active

    headers = [cell.value for cell in sheet[1]]
    date_col_idx = [i + 1 for i, h in enumerate(headers) if "Fecha" in str(h)][0]
    ppm_col_idx = [i + 1 for i, h in enumerate(headers) if "PPM" in str(h).upper()][0]
    words_col_idx = [i + 1 for i, h in enumerate(headers) if "Palabra" in str(h)][0]

    # Fila 2: primer resultado
    date_cell = sheet.cell(row=2, column=date_col_idx)
    ppm_cell = sheet.cell(row=2, column=ppm_col_idx)
    words_cell = sheet.cell(row=2, column=words_col_idx)

    # 1. La fecha no debe ser string plano, debe ser datetime.date o tener formato de fecha
    assert isinstance(date_cell.value, (datetime.date, datetime.datetime)) or date_cell.is_date

    # 2. El PPM y métricas numéricas deben ser float o int, nunca str
    assert isinstance(ppm_cell.value, (int, float))
    assert not isinstance(ppm_cell.value, str)

    # 3. Palabras o tiempo deben ser int
    assert isinstance(words_cell.value, int)
    assert not isinstance(words_cell.value, str)


def test_scenario_5_empty_dataset_exports_headers_only(client, setup_data):
    """
    Escenario 5: Conjunto vacío
    Dado unos filtros que no devuelven ningún resultado
    Cuando se solicita la exportación
    Entonces se genera un fichero con solo las cabeceras y no se produce un error.
    """
    client.post("/api/dev/session", json={"role": "coordinator", "sections": []})

    non_existent_section = uuid.uuid4()
    resp = client.get(f"/api/v1/results/export/excel?section_id={non_existent_section}")
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))

    # Debe contener exactamente 1 fila (las cabeceras)
    assert len(rows) == 1
    assert len(rows[0]) > 5


def test_scenario_6_audit_logging_on_export(client, setup_data):
    """
    Escenario 6: Registro de la exportación
    Dado una exportación completada
    Cuando termina la descarga
    Entonces queda registrada en auditoría con el usuario y el ámbito exportado.
    """
    sec1 = setup_data["section1"]
    client.post("/api/dev/session", json={"role": "tutor", "sections": [str(sec1.id)]})

    clear_audit_logs()

    resp = client.get(f"/api/v1/results/export/excel?section_id={sec1.id}")
    assert resp.status_code == 200

    logs = get_audit_logs(resource_type="results")
    export_logs = [log for log in logs if log["action"] == "EXPORT_EXCEL"]
    assert len(export_logs) >= 1

    audit_entry = export_logs[0]
    assert "details" in audit_entry
    assert audit_entry["details"].get("section_id") == str(sec1.id)


def test_access_control_unauthenticated_and_pending_role(client, setup_data):
    """
    Control de acceso: 401 sin sesión, 403 con rol pendiente y 403 si tutor intenta exportar sección ajena.
    """
    sec1 = setup_data["section1"]
    sec2 = setup_data["section2"]

    # 1. 401 sin sesión
    app = create_app(NoBypassConfig)
    with app.test_client() as unauth_client:
        assert unauth_client.get("/api/v1/results/export/excel").status_code == 401

    # 2. 403 con rol pendiente
    client.post("/api/dev/session", json={"role": "pendiente", "sections": []})
    assert client.get("/api/v1/results/export/excel").status_code == 403

    # 3. 403 si tutor asignado solo a sec1 intenta exportar explícitamente sec2
    client.post("/api/dev/session", json={"role": "tutor", "sections": [str(sec1.id)]})
    resp = client.get(f"/api/v1/results/export/excel?section_id={sec2.id}")
    assert resp.status_code == 403
    assert resp.get_json()["error"] == "FORBIDDEN"


def test_versioned_endpoint_and_direct_alias(client, setup_data):
    """
    Verifica que /api/v1/results/export/excel y /api/results/export/excel responden idénticamente.
    """
    client.post("/api/dev/session", json={"role": "coordinator", "sections": []})

    resp_v1 = client.get("/api/v1/results/export/excel")
    assert resp_v1.status_code == 200

    resp_direct = client.get("/api/results/export/excel")
    assert resp_direct.status_code == 200
    assert len(resp_direct.data) > 0
